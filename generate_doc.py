import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = Document()

# Set page margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

# Colors
C_RED = RGBColor(211, 41, 32)       # #d32920
C_INK = RGBColor(24, 23, 21)        # #181715
C_MUTED = RGBColor(119, 115, 108)   # #77736c

# 1. Header tag
p_tag = doc.add_paragraph()
r_tag = p_tag.add_run("MANIFESTE DE LANCEMENT · ÉDUCATION POPULAIRE · ACTION DIRECTE")
r_tag.font.name = "Arial"
r_tag.font.size = Pt(9.5)
r_tag.font.bold = True
r_tag.font.color.rgb = C_RED

# 2. Main Title
p_title = doc.add_paragraph()
r_title = p_title.add_run("Bienvenue sur Jeunesse Libertaire :\nNos voix, nos luttes, notre média !")
r_title.font.name = "Arial"
r_title.font.size = Pt(24)
r_title.font.bold = True
r_title.font.color.rgb = C_INK
p_title.paragraph_format.space_after = Pt(12)

# 3. Chapeau (Deck)
p_deck = doc.add_paragraph()
r_deck = p_deck.add_run(
    "Face au matraquage idéologique, à la précarité généralisée et au renforcement autoritaire de l'État, "
    "la jeunesse refuse le silence et la résignation. Bienvenue sur Jeunesse Libertaire, un média autogéré, "
    "de combat et sans frontières, conçu par et pour les lycéen·nes, les étudiant·es et les jeunes travailleur·euses."
)
r_deck.font.name = "Georgia"
r_deck.font.size = Pt(12)
r_deck.font.italic = True
r_deck.font.color.rgb = C_INK
p_deck.paragraph_format.space_after = Pt(18)

# Helper function for headings
def add_h2(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = C_RED
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)

def add_body(text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.font.name = "Georgia"
        rb.font.size = Pt(10.5)
        rb.font.bold = True
        rb.font.color.rgb = C_INK
    r = p.add_run(text)
    r.font.name = "Georgia"
    r.font.size = Pt(10.5)
    r.font.color.rgb = C_INK
    p.paragraph_format.space_after = Pt(6)

# Section 1
add_h2("1. Pourquoi Jeunesse Libertaire ?")
add_body(
    "Partout, les médias dominants et les institutions prétendent parler à notre place. Ils nous infantilisent, "
    "caricaturent nos colères et cherchent à étouffer la moindre contestation pour préserver l'ordre établi. "
    "Dans nos lycées, nos facs et nos ateliers, la réalité quotidienne est brutale : sélection sociale accrue, "
    "tri sélectif dès le plus jeune âge, précarité qui explose et présence policière omniprésente dès que nous relevons la tête."
)
add_body(
    "Nous n'avons rien à attendre de leurs réformes, de leurs promesses électorales ni de leurs négociations de couloir. "
    "L'émancipation ne se mendie pas auprès des puissants : elle s'arrache et s'organise à la base. "
    "Pour transformer le monde, nous devons d'abord nous réapproprier nos propres armes d'information, de débat et de lutte collective."
)

# Section 2
add_h2("2. À quoi sert ce média ?")
add_body(
    " Donner un écho direct aux blocages, aux assemblées générales, aux cortèges spontanés et aux grèves. "
    "Raconter sans filtre ce qui se passe sur le terrain, nos victoires tactiques comme nos apprentissages.",
    bold_prefix="• Rendre visibles les luttes du terrain :"
)
add_body(
    " Sortir du bourrage de crâne officiel en partageant des synthèses historiques, des analyses percutantes "
    "sur le capitalisme et l'État, et des guides concrets d'auto-défense (droits face à la police, caisses de grève, gestion des AG).",
    bold_prefix="• Auto-formation & Éducation populaire :"
)
add_body(
    " Briser l'isolement entre les bahuts, les villes et les universités. Offrir une tribune ouverte "
    "à toutes celles et ceux qui luttent, sans intermédiaire bureaucratique.",
    bold_prefix="• Coopération horizontale & Autonomie :"
)

# Section 3
add_h2("3. Comment prendre part à l'aventure ?")
add_body(
    "Ce média n'appartient à aucun parti politique, à aucun patron de presse ni à aucune structure corporatiste. "
    "Il vit exclusivement de l'énergie, de la plume et de la solidarité de la jeunesse révoltée."
)
add_body(
    "1. Proposer un texte, une analyse ou un témoignage : rendez-vous sur le formulaire public du site. "
    "Chaque proposition est relue et discutée collectivement par les compagnes et compagnons du collectif."
)
add_body(
    "2. Débattre et enrichir : réagissez sous les publications dans les espaces de discussion modérés."
)
add_body(
    "3. Diffuser et coller : imprimez les tracts et articles, partagez les vidéos sur TikTok et Instagram, "
    "faites tourner le site dans vos assemblées générales !"
)

# Conclusion Call to Action
p_box = doc.add_paragraph()
r_box = p_box.add_run("\nREFUSER LA RÉSIGNATION · CONSTRUIRE L'ÉMANCIPATION · SOLIDARITÉ ET ACTION DIRECTE !")
r_box.font.name = "Arial"
r_box.font.size = Pt(11)
r_box.font.bold = True
r_box.font.color.rgb = C_RED
p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_box.paragraph_format.space_before = Pt(18)
p_box.paragraph_format.space_after = Pt(12)

# Footer note
p_foot = doc.add_paragraph()
r_foot = p_foot.add_run("Jeunesse Libertaire — Média libre, autogéré et sans frontières · Licence Creative Commons CC BY-SA 4.0")
r_foot.font.name = "Arial"
r_foot.font.size = Pt(8.5)
r_foot.font.color.rgb = C_MUTED
p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.save("/home/anar/jeunesse-libertaire/assets/bienvenue-article.docx")
print("DOCX_GENERATED_SUCCESSFULLY")
